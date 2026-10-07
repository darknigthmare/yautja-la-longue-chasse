"""Extract original V6 source cells read-only; never execute or edit the XLSX."""
from pathlib import Path
import hashlib
import json
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
ORIGINAL = ROOT / "work-local/drive-import-20261007/downloads/Yautja_V6_Bible_Cumulative_et_Dialogues.xlsx"
EXPECTED = "87ad5d1a6eaf49f2908a4457ef3f97b8f34db481a24bbc121d2e924559b19183"
assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() == EXPECTED
book = openpyxl.load_workbook(ORIGINAL, read_only=True, data_only=False)
specs = {
    "Campagne commune": ([27], 11),
    "Vaisseau personnel": (list(range(6, 21, 2)), 12),
    "Scènes de dialogue V6": ([598], 17),
    "Choix et actions V6": ([1218, 1219], 9),
}
sheets = []
for name, (numbers, width) in specs.items():
    wanted = set(numbers)
    rows = []
    for number, cells in enumerate(book[name].iter_rows(max_row=max(numbers), max_col=width), 1):
        if number not in wanted:
            continue
        rows.append({"number": number, "cells": [
            {"address": cell.coordinate, "value": cell.value}
            for cell in cells if cell.value is not None
        ]})
    assert len(rows) == len(numbers), name
    sheets.append({"name": name, "rows": rows})
result = {"schemaVersion": 1, "source": {"workbook": ORIGINAL.name,
    "sha256": EXPECTED, "totalWorkbookSheets": len(book.sheetnames)}, "sheets": sheets}
book.close()
target = ROOT / "app/game/data/shipPreparationSourceV88.json"
target.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"Extracted {sum(len(s['rows']) for s in sheets)} source rows from {len(sheets)} sheets")
