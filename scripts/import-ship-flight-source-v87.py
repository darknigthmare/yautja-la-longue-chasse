"""Read-only extraction: never execute workbook content or change the XLSX."""
from pathlib import Path
import hashlib
import json
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
ORIGINAL = ROOT / "work-local/drive-import-20261007/downloads/Yautja_V6_Bible_Cumulative_et_Dialogues.xlsx"
EXPECTED = "87ad5d1a6eaf49f2908a4457ef3f97b8f34db481a24bbc121d2e924559b19183"
assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() == EXPECTED, "Unexpected source workbook"
book = openpyxl.load_workbook(ORIGINAL, read_only=True, data_only=False)
specs = {"V5 Manœuvres spatiales": ([6, 7, 8, 9, 16, 28], 12), "Vaisseau personnel": ([20], 12)}
sheets = []
for name, (numbers, width) in specs.items():
    sheet = book[name]
    rows = []
    for number in numbers:
        cells = [{"address": sheet.cell(number, column).coordinate, "value": sheet.cell(number, column).value}
                 for column in range(1, width + 1) if sheet.cell(number, column).value is not None]
        rows.append({"number": number, "cells": cells})
    sheets.append({"name": name, "rows": rows})
result = {"schemaVersion": 1, "source": {"workbook": ORIGINAL.name, "sha256": EXPECTED, "totalWorkbookSheets": len(book.sheetnames)}, "sheets": sheets}
book.close()
target = ROOT / "app/game/data/shipFlightSourceV87.json"
target.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"Extracted {sum(len(s['rows']) for s in sheets)} source rows from {len(sheets)} sheets")
