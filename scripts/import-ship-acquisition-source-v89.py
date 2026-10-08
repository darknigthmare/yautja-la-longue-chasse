#!/usr/bin/env python3
"""Read-only V6 acquisition source extraction; values are data, never commands.

Run only after disk space is available. This script does not save/modify the
workbook or evaluate formulas. Output retains original source cells/provenance.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "work-local/drive-import-20261007/downloads/Yautja_V6_Bible_Cumulative_et_Dialogues.xlsx"
DEFAULT_OUTPUT = ROOT / "app/game/data/shipAcquisitionSourceV89.json"
SOURCE_SHA = "87ad5d1a6eaf49f2908a4457ef3f97b8f34db481a24bbc121d2e924559b19183"
RANGES = (
    ("Campagne commune", 25, 28, 11),
    ("Scènes de dialogue V6", 597, 599, 17),
    ("Répliques V6", 4923, 4946, 11),
    ("Choix et actions V6", 1216, 1221, 9),
    ("Voix V6", 208, 208, 9),
    ("V5 États communs", 22, 25, 10),
)


def extract(source_path: Path) -> dict:
    digest = hashlib.sha256(source_path.read_bytes()).hexdigest()
    if digest != SOURCE_SHA:
        raise ValueError("Workbook SHA differs from reviewed V6 source; no output written")
    workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=False)
    try:
        if len(workbook.sheetnames) != 139:
            raise ValueError("Reviewed workbook must contain 139 original worksheets")
        sheets = []
        for name, first, last, width in RANGES:
            sheet = workbook[name]
            rows = []
            for row_number, cells in enumerate(sheet.iter_rows(min_row=first, max_row=last, min_col=1, max_col=width), first):
                rows.append({
                    "number": row_number,
                    "cells": [{"address": cell.coordinate, "value": cell.value,
                               **({"formula": cell.value} if cell.data_type == "f" else {})}
                              for cell in cells if cell.value is not None],
                })
            sheets.append({"name": name, "rows": rows})
        return {
            "schemaVersion": 1,
            "source": {"workbook": source_path.name, "sha256": digest, "totalWorkbookSheets": 139},
            "scope": "Original R2-M020/021/022/023 source subset; no acquisition, economics, fuel or voice clip is generated.",
            "sheets": sheets,
        }
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    data = extract(args.source)
    serialized = json.dumps(data, ensure_ascii=False, indent=2) + "\n"
    # Same-directory replacement avoids a partial JSON output on a failed write.
    temporary = args.output.with_name(args.output.name + ".new")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary.write_text(serialized, encoding="utf-8")
    temporary.replace(args.output)
    count = sum(len(sheet["rows"]) for sheet in data["sheets"])
    print(f"Extracted {count} original rows; workbook unchanged")


if __name__ == "__main__":
    main()
