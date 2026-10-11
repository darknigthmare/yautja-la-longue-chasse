"""Read-only extraction. Never executes workbook content or changes the XLSX."""
from pathlib import Path
import hashlib
import json
import openpyxl

SOURCE = Path('work-local/drive-import-20261007/downloads/Yautja_V6_Bible_Cumulative_et_Dialogues.xlsx')
SHA = '87ad5d1a6eaf49f2908a4457ef3f97b8f34db481a24bbc121d2e924559b19183'
if hashlib.sha256(SOURCE.read_bytes()).hexdigest() != SHA:
    raise SystemExit('Refusing a different workbook revision')
book = openpyxl.load_workbook(SOURCE, read_only=True, data_only=False)
sheets = []
for name, rows, columns in [('Campagne commune', [21, 22, 23, 24, 25], range(1, 10)),
                            ('Scènes de dialogue V6', [594, 595, 596], range(1, 12))]:
    sheet = book[name]
    sheets.append({'name': name, 'rows': [{'number': row, 'cells': [
        {'address': sheet.cell(row, column).coordinate, 'value': sheet.cell(row, column).value}
        for column in columns]} for row in rows]})
result = {'source': {'workbook': SOURCE.name, 'sha256': SHA},
          'runtimeScope': 'training-only-not-R2-M018-completion', 'sheets': sheets}
Path('app/game/data/campaignTrainingSourceV90.json').write_text(
    json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
